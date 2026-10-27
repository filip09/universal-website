from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1C1D2E")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
ID_FONT = Font(name=FONT, size=10, bold=True)
CODE = Font(name="Courier New", size=10)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="7A7890")
NOTE_FILL = PatternFill("solid", fgColor="F3F2F7")
CONTENT_FILL = PatternFill("solid", fgColor="FFF8E6")
LINE = Border(bottom=Side(style="thin", color="E2E0EA"))
TOP = Alignment(vertical="top", wrap_text=True)

wb = Workbook()


def sheet(name, cols, rows, id_cols=(0,), note_cols=(), code_cols=(), validations=None, tab=None):
    ws = wb.create_sheet(name)
    for c, (head, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 24
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment, cell.border = TOP, LINE
            cell.font = ID_FONT if (c - 1) in id_cols else BODY
            if (c - 1) in code_cols:
                cell.font, cell.fill = CODE, CONTENT_FILL
            if (c - 1) in note_cols:
                cell.font, cell.fill = NOTE_FONT, NOTE_FILL
    for c in note_cols:
        for r in range(len(rows) + 2, 200):
            ws.cell(row=r, column=c + 1).fill = NOTE_FILL
    for c in code_cols:
        for r in range(len(rows) + 2, 200):
            ws.cell(row=r, column=c + 1).fill = CONTENT_FILL
    ws.freeze_panes = "B2" if len(cols) > 3 else "A2"
    for col_idx, options in (validations or {}).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Vyberte jednu z nabízených hodnot.", "Neznámá hodnota"
        ws.add_data_validation(dv)
        letter = ws.cell(row=1, column=col_idx + 1).column_letter
        dv.add(f"{letter}2:{letter}500")
    if tab:
        ws.sheet_properties.tabColor = tab
    return ws


YESNO = ["yes", "no"]

# ---------- pages: the heart of it ----------
pages = [
    ("intro", 10, "Úvod", "Web, kterým se listuje", 265, "yes",
     "text: t-intro\n"
     "text: t-sheet\n"
     "mezititulek: Zkuste ho přerušit\n"
     "text: Klikněte na záložku a hned na další, než se stránka dootočí. Stoh **změní směr** uprostřed pohybu.\n"
     "text: t-contact | centered",
     "Jak funguje web, kterým se listuje jako knihou.", "img-hero",
     "Kombinace id z texts a textu psaného rovnou do buňky."),
    ("turn", 20, "Otočení", "Tři fáze otočení", 198, "yes",
     "text: t-turn\n"
     "seznam: turn-beats | steps\n"
     "text: t-fast",
     "Zvednutí, profil, dosednutí: co se děje mezi stránkami.", "", ""),
    ("gallery", 30, "Galerie", "Obrázky na skle", 328, "yes",
     "obrázek: img-hero | full\n"
     "galerie: gal-evening | grid\n"
     "text: Obrázky se načítají, až když se jejich stránka blíží dopředu.",
     "", "img-hero", ""),
    ("blocks", 40, "Bloky", "Bloky z tabulky", 158, "yes",
     "video: https://youtu.be/aqz-KE-bpKQ | wide\n"
     "vložit: html-clock\n"
     "poznámka: Video se zastaví, když se stránka otočí pryč.",
     "", "", "Odkaz na video rovnou v buňce, bez řádku v media."),
    ("next", 50, "Dál", "Co postavíme dál", 32, "yes",
     "úvod: Až bude pohyb sedět, přijde platforma pod ním.\n"
     "seznam: next-steps | steps\n"
     "oddělovač:\n"
     "text: t-contact",
     "", "", "Dokud není hotová, dejte visible = no."),
]
sheet("pages", [("id", 12), ("order", 7), ("tab_label", 12), ("title", 26), ("hue", 6), ("visible", 8),
                ("content", 64), ("description", 36), ("share_image", 13), ("notes", 30)],
      pages, code_cols=(6,), note_cols=(9,), validations={5: YESNO}, tab="F0A33C")

# ---------- texts ----------
sheet("texts", [("id", 14), ("content", 80), ("notes", 34)], [
    ("t-intro", "úvod: Každá stránka se čte jako běžná webová stránka a roluje, jak dlouho potřebuje. Rozdíl je ve chvíli **mezi** stránkami.\nnadpis: Jedna stránka v čase\ntext: Zatímco čtete, ostatní stránky nejsou vidět. Že existují, prozrazují jen záložky.\n\nKaždá stránka je list tónovaného skla. Barvy pozadí prosvítají skrz.", "Prázdný řádek uvnitř text = nový odstavec."),
    ("t-sheet", "nadpis: Postaveno z tabulky\ntext: Stránky, jejich pořadí, barva i veškerý obsah pocházejí z *Google Sheets*.\npoznámka: Tabulka se čte při sestavení webu, nikdy ne ve chvíli, kdy návštěvník čeká.", ""),
    ("t-turn", "úvod: Otočení trvá necelou sekundu. Pohání ho pružiny, takže má váhu a na konci lehce dosedne.", ""),
    ("t-fast", "nadpis: Proč je to rychlé\ntext: Pohyb používá jen 3D transformace a průhlednost.\ncitace: Ať to působí jako otáčení skutečné stránky, jen lehčí.\nautor: Studio\nseznam: žádné přepočítávání rozvržení; vzdálené stránky skryté; malé plátno na pozadí", "Citace + autor a seznam na jednom řádku."),
    ("t-contact", "mezititulek: Chcete takový web?\ntext: Napište nám a uděláme z něj ten váš.\ntlačítko: Napište nám -> https://www.example.com/kontakt", "Použito na dvou stránkách: intro a next."),
], note_cols=(2,), code_cols=(1,), tab="D9548C")

# ---------- media ----------
media_rows = [
    ("img-hero", "image", "hero.jpg", "Růžové a fialové světlo nad tmavou plochou", "Ukázkový obrázek", "", "center", "Jen název souboru ze složky media."),
    ("img-favicon", "image", "ikona.png", "Ikona webu", "", "", "center", "Čtvercový obrázek, aspoň 512 × 512 px."),
    ("vid-ukazka", "video", "https://youtu.be/aqz-KE-bpKQ", "Přehrát ukázkové video", "Ukázkové video", "", "", "Video sem dávejte, jen když chcete popisek. Jinak stačí odkaz přímo v pages."),
]
for n, alt in enumerate(["Tyrkysové a modré světlo", "Zlaté a korálové světlo", "Fialové a azurové světlo",
                         "Oranžové a krémové světlo", "Zelené a modré světlo", "Purpurové a indigové světlo"], 1):
    media_rows.append((f"img-ev-{n:02d}", "image", f"galerie/vecer-{n:02d}.jpg", alt, "", "", "center", ""))
sheet("media", [("id", 14), ("kind", 8), ("source", 46), ("alt", 34), ("caption", 28), ("credit", 14), ("focus", 9), ("notes", 44)],
      media_rows, note_cols=(7,), validations={1: ["image", "video"], 6: ["center", "top", "bottom", "left", "right"]}, tab="8A63D2")

# ---------- items ----------
items = [
    ("turn-beats", 10, "Zvednutí", "Přední stránka zprůhlední a ustoupí. Stránky za ní prosvitnou sklem.", "", "", "yes"),
    ("turn-beats", 20, "Profil", "Celý stoh se natočí bokem a posune, dokud zvolená stránka není vepředu.", "", "", "yes"),
    ("turn-beats", 30, "Dosednutí", "Nová stránka znovu zprůhlední do plna a stoh se vrátí rovně.", "", "", "yes"),
]
for n in range(1, 7):
    items.append(("gal-evening", n * 10, "", "", f"img-ev-{n:02d}", "", "yes"))
for n, (t, x) in enumerate([
    ("Doladit otočení", "Hloubka, úhel a rychlost, dokud to nesedí na počítači i v telefonu."),
    ("Navrhnout tabulku", "Listy pro stránky, texty, média a styly."),
    ("Sestavení", "Skript přečte tabulku a vytvoří statické stránky."),
    ("Hosting", "Statické soubory na Cloudflare Pages nebo Netlify."),
    ("Bloky", "Text, obrázek, galerie, video a vložené HTML."),
], 1):
    items.append(("next-steps", n * 10, t, x, "", "", "yes"))
sheet("items", [("group", 14), ("order", 7), ("title", 20), ("text", 60), ("image", 14), ("link", 30), ("visible", 8)],
      items, validations={6: YESNO}, tab="2BB3A0")

# ---------- embeds ----------
sheet("embeds", [("id", 14), ("html", 70), ("height", 8), ("notes", 40)], [
    ("html-clock", "<div id=\"clock\" style=\"font-size:42px\"></div>\n<script>setInterval(()=>clock.textContent=new Date().toLocaleTimeString('cs-CZ'),1000)</script>", 80, "Běží jen na stránce, která je vepředu. Výška v px, prázdné = automaticky."),
], note_cols=(3,), code_cols=(1,), tab="8A63D2")

# ---------- site ----------
sheet("site", [("setting", 22), ("value", 46), ("co to dělá", 58)], [
    ("site_name", "Universal Website", "Název v záložce prohlížeče a při sdílení odkazu."),
    ("languages", "cs", "Jazyky webu, první je hlavní: cs, nebo cs, en. Další jazyk píšete do sloupců s koncovkou _en (viz Návod › Více jazyků)."),
    ("domain", "www.example.com", "Kde web běží. Pro sdílené odkazy a mapu webu."),
    ("home_page", "intro", "Id stránky, která se otevře jako první."),
    ("description", "Web, kterým se listuje, postavený z tabulky.", "Výchozí text pro vyhledávače a náhledy odkazů."),
    ("share_image", "img-hero", "Výchozí obrázek při sdílení odkazu (id z media)."),
    ("favicon", "img-favicon", "Malá ikona v záložce prohlížeče (id z media)."),
    ("footer", "drobné: © {rok} {web}", "Patička na každé stránce. Stejný zápis jako content. {rok} je vždy aktuální rok, {web} je site_name. Prázdné = © rok a název webu."),
    ("form_email", "", "Kam chodí zprávy z formuláře (formulář: kontakt). Na webu se nezobrazí. Více adres oddělte čárkou."),
    ("address", "", "Adresa pro mapu (mapa:), např. Vinohradská 12, Praha 2. Na webu se zobrazí pod mapou."),
    ("map_zoom", "15", "Přiblížení mapy: 1 = celý svět, 15 = ulice, 18 = dům."),
], note_cols=(2,), tab="6B5BFF")

# ---------- styles ----------
PROFILES = ["manual", "terakota", "jantar", "koral", "broskev", "karamel", "ruze", "vinova", "pisek",
            "seda", "tron", "soumrak", "zeme", "pastel", "sever"]
styles_ws = sheet("styles", [("setting", 24), ("value", 16), ("možnosti", 26), ("co to dělá", 58)], [
    ("color_profile", "manual", "roletka (viz Návod)", "Barevný profil webu. manual = barvy podle sloupce hue v pages a nastavení níže. Ostatní profily nastaví barvy samy."),
    ("background", "drift", "drift / still / solid", "drift: barvy se pomalu pohybují. still: stejné barvy bez pohybu. solid: jedna barva."),
    ("background_speed", "slow", "slow / normal", "Jak rychle se barvy pozadí mění."),
    ("background_follows_page", "yes", "yes / no", "Pozadí přebírá barvu stránky, která je vepředu."),
    ("glass_opacity", 0.6, "0.3 – 0.95", "Jak plné jsou stránky. Méně = víc pozadí prosvítá."),
    ("page_width", 860, "640 – 1200 (px)", "Největší šířka stránky na počítači."),
    ("page_radius", 22, "0 – 40 (px)", "Zaoblení rohů stránky."),
    ("font_display", "Gloock", "libovolný Google Font", "Písmo nadpisů."),
    ("font_body", "Figtree", "libovolný Google Font", "Písmo textu."),
    ("font_utility", "JetBrains Mono", "libovolný Google Font", "Písmo malých popisků a čísel."),
    ("text_size", 16, "14 – 20 (px)", "Základní velikost textu. Nadpisy se odvodí."),
    ("text_color", "#ECE9F4", "hex barva", "Hlavní barva textu."),
    ("muted_color", "#A29FB7", "hex barva", "Barva popisků a vedlejšího textu."),
    ("hue_saturation", 70, "0 – 100", "Jak syté jsou barvy stránek. 0 = jen šedé."),
    ("turn_speed", "normal", "slow / normal / quick", "Rychlost otočení."),
    ("turn_angle", 30, "0 – 45 (stupňů)", "Jak moc se stoh natočí bokem. 0 natočení vypne."),
    ("stack_depth", 5, "1 – 8", "Kolik stránek je při otáčení vidět za tou přední."),
    ("tabs_desktop", "right", "right / left", "Na které straně jsou záložky na počítači."),
], note_cols=(2, 3), tab="6B5BFF")
_dv = DataValidation(type="list", formula1='"' + ",".join(PROFILES) + '"', allow_blank=False)
_dv.error, _dv.errorTitle = "Vyberte jeden z nabízených profilů.", "Neznámý profil"
styles_ws.add_data_validation(_dv)
_dv.add("B2")

# ---------- Guide (Czech) ----------
g = wb.active
g.title = "Návod"
g.sheet_properties.tabColor = "1C1D2E"
for col, w in zip("ABCD", (22, 18, 44, 58)):
    g.column_dimensions[col].width = w
r = 1


LANG_LINES = [
    "V listu site napište do languages všechny jazyky, první je hlavní: cs, en. Hlavní jazyk je na adrese webu, další pod /en/.",
    "Překlad píšete do sloupců s koncovkou jazyka vedle původních: tab_label_en, title_en, content_en, description_en v pages, content_en v texts, title_en a text_en v items, alt_en v media, html_en v embeds.",
    "id_en v pages je adresa stránky v angličtině, např. why-us → /en/why-us/. Prázdné = stejné id jako v češtině.",
    "Odkazy v překladu můžete psát s českým i anglickým id (/proc-s-nami/ i /why-us/), vždy povedou na anglickou verzi stránky.",
    "visible_en: no skryje stránku nebo položku jen v angličtině.",
    "Obrázky, barvy, pořadí a nastavení vzhledu jsou společné. V content_en stačí napsat stejné obrázky a galerie.",
    "Chybí-li překlad, zobrazí se původní text a v listu stav se objeví upozornění.",
    "Návštěvník přepíná jazyk u záložek. Má-li prohlížeč nastavený jiný jazyk webu, otevře se mu napoprvé rovnou ten. Jeho volba se pamatuje.",
    "V site můžete přeložit i site_name_en, description_en a footer_en.",
]


def put(row, vals, font=BODY, fill=None):
    for c, v in enumerate(vals, 1):
        cell = g.cell(row=row, column=c, value=v)
        cell.font, cell.alignment = font, TOP
        if fill:
            cell.fill = fill


def section(title):
    global r
    r += 1
    g.cell(row=r, column=1, value=title).font = Font(name=FONT, size=13, bold=True)
    r += 1


def table(head, rows, code_first=False):
    global r
    put(r, head, HEAD_FONT, HEAD_FILL)
    r += 1
    for row in rows:
        put(r, row)
        for c in range(1, len(head) + 1):
            g.cell(row=r, column=c).border = LINE
        if code_first:
            g.cell(row=r, column=1).font = CODE
            g.cell(row=r, column=4).font = CODE
        r += 1


def para(text, height=30):
    global r
    put(r, [text])
    g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    g.row_dimensions[r].height = height
    r += 1


g.cell(row=r, column=1, value="Universal Website · návod k tabulce").font = Font(name=FONT, size=18, bold=True)
r += 1
para("Stránku skládáte přímo v listu pages, ve sloupci content (žluté buňky). Každý řádek v buňce je jeden kus obsahu, v pořadí shora dolů. Ukázkové řádky odpovídají prototypu, přepište je nebo smažte.")

section("Jak to do sebe zapadá")
table(["list", "jeden řádek je", "hlavní sloupce", "k čemu"], [
    ("pages", "jedna stránka = jedna záložka", "id, order, tab_label, title, hue, content", "Tady skládáte. Obsah stránky píšete do content."),
    ("texts", "jeden znovupoužitelný text", "id, content", "Delší texty, které chcete mít zvlášť nebo na více stránkách."),
    ("media", "jeden obrázek nebo video", "id, source, alt", "Volitelné: obrázky, kterým chcete dát popis nebo popisek. Jinak stačí napsat název souboru přímo v pages."),
    ("items", "jedna karta, krok nebo fotka galerie", "group, order", "Seznamy karet a galerie: řádky se stejnou group patří k sobě."),
    ("embeds", "kus HTML", "id, html", "Widgety a cokoli vlastního."),
    ("site", "jedno nastavení", "setting, value", "Název, doména, úvodní stránka, patička."),
    ("styles", "jedno nastavení vzhledu", "setting, value", "Písma, barvy, sklo, rychlost a úhel otočení."),
])

section("Skládání stránky: sloupec content")
para("Každý kus obsahu začíná klíčovým slovem a dvojtečkou. Za ním je buď id z jiného listu, odkaz, nebo rovnou text. Nový řádek v buňce: Ctrl+Enter (Windows) nebo ⌘+Enter (Mac). Česká i anglická slova fungují stejně.", 42)
table(["klíčové slovo", "anglicky", "co za ním může být", "příklad"], [
    ("text:", "text:", "id z texts, nebo rovnou text", "text: t-intro"),
    ("obrázek:", "image:", "id z media, název souboru ze složky media, nebo odkaz", "obrázek: hero.jpg | full"),
    ("galerie:", "gallery:", "group z items", "galerie: gal-evening | grid"),
    ("video:", "video:", "odkaz na YouTube / Vimeo, nebo id z media", "video: https://youtu.be/abc123 | wide"),
    ("seznam:", "list:", "group z items (karty), nebo položky oddělené ;", "seznam: turn-beats | steps"),
    ("vložit:", "embed:", "id z embeds", "vložit: html-clock"),
    ("formulář:", "form:", "kontakt (zprávy chodí na form_email z listu site)", "formulář: kontakt"),
], code_first=True)
r += 1
para("Jedno slovo bez mezer, které odpovídá existujícímu id, se vždy vezme jako odkaz na ten obsah. Cokoli jiného se zobrazí jako text.")
para("Za svislou čáru | můžete napsat styl. Bez stylu se použije výchozí (první v tabulce níže).")
table(["obsah", "", "styly", "co znamenají"], [
    ("text", "", "normal / wide / centered / large", "large = větší písmo pro výrazné sdělení"),
    ("obrázek", "", "inset / full", "full = přes celou šířku stránky"),
    ("galerie", "", "grid / row", "row = řada, která se posouvá do strany"),
    ("video", "", "normal / wide", ""),
    ("seznam", "", "steps / cards / bullets / faq", "steps čísluje karty podle order; cards bez čísel; faq = rozbalovací otázky (title = otázka, text = odpověď)"),
])

section("Formátování textu")
para("Tato slova můžete psát přímo v content stránky i v listu texts. Vše v jedné buňce, každou část na nový řádek.")
table(["klíčové slovo", "anglicky", "co vytvoří", "příklad"], [
    ("úvod:", "lead:", "Velký úvodní odstavec", "úvod: Každá stránka se čte jako běžný web."),
    ("text:", "text:", "Odstavec. Prázdný řádek uvnitř = nový odstavec.", "text: Ostatní stránky nejsou vidět."),
    ("nadpis:", "heading:", "Nadpis oddílu", "nadpis: Proč je to rychlé"),
    ("mezititulek:", "subheading:", "Menší nadpis", "mezititulek: Zkuste ho přerušit"),
    ("citace:", "quote:", "Zvýrazněná citace", "citace: Ať to působí jako skutečná stránka."),
    ("autor:", "author:", "Kdo to řekl; patří k citaci nad ním", "autor: Studio"),
    ("seznam:", "list:", "Odrážky oddělené ; nebo řádky začínající -", "seznam: rychlé; lehké; z tabulky"),
    ("poznámka:", "note:", "Rámeček s poznámkou", "poznámka: Tabulka se čte při sestavení."),
    ("tlačítko:", "button:", "Tlačítko: popis -> odkaz", "tlačítko: Napište nám -> https://example.com"),
    ("oddělovač:", "divider:", "Tenká čára; za dvojtečkou nic", "oddělovač:"),
    ("mapa:", "map:", "Mapa Google. Prázdné = adresa z listu site, nebo jiná adresa; | light = světlá", "mapa:"),
    ("mezera:", "space:", "Volné místo. Prázdné = běžná, nebo velká / malá", "mezera: velká"),
    ("drobné:", "small:", "Drobný text", "drobné: © {rok} {web}"),
], code_first=True)
r += 1
table(["uvnitř textu", "", "výsledek", "příklad"], [
    ("**slova**", "", "tučně", "stoh **změní směr**"),
    ("*slova*", "", "kurzíva", "z *Google Sheets*"),
    ("[slova](odkaz)", "", "odkaz", "viz [naše studio](https://example.com)"),
], code_first=True)

section("Celá stránka v jedné buňce")
ex = ("úvod: Až bude pohyb sedět, přijde platforma pod ním.\n"
      "text: t-intro\n"
      "obrázek: img-hero | full\n"
      "mezititulek: Jak to funguje\n"
      "seznam: turn-beats | steps\n"
      "video: https://youtu.be/abc123\n"
      "citace: Ať to působí jako skutečná stránka.\n"
      "autor: Studio\n"
      "galerie: gal-evening\n"
      "tlačítko: Napište nám -> https://example.com/kontakt")
c = g.cell(row=r, column=1, value=ex)
c.font, c.alignment, c.fill = CODE, TOP, CONTENT_FILL
g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
g.row_dimensions[r].height = 150
r += 1

section("Barevné profily (styles › color_profile)")
para("V listu styles vyberte v řádku color_profile jeden profil. Pozadí zůstává živé a jemně proměnlivé, profil mu dává základní tón. Kromě manual profily nepoužívají sloupec hue ani barvy textu z listu styles.", 42)
table(["profil", "", "jak působí", "barvy"], [
    ("manual", "", "Barvu každé stránky určuje sloupec hue v listu pages.", "podle vás"),
    ("terakota", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Teplá cihlová.", "#C8643B"),
    ("jantar", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Medová, zlatavá.", "#D9932E"),
    ("koral", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Svěží korálová.", "#E8715A"),
    ("broskev", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Jemná broskvová.", "#E89B7A"),
    ("karamel", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Karamelově hnědá.", "#B9773F"),
    ("ruze", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Prášková růže.", "#C97B84"),
    ("vinova", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Hluboká vínová.", "#9E3B5A"),
    ("pisek", "", "Jedna barva v odstínech, pozadí v doplňkové barvě. Světlý písek.", "#C9A77A"),
    ("seda", "", "Stupně šedi, bez barvy. Pozadí v tmavě šedých tónech.", "#9C9C9E"),
    ("tron", "", "Neonově azurová s oranžovou, záře na hranách, pozadí v modrých a oranžových tónech.", "#00E5FF · #FF7A1A"),
    ("soumrak", "", "Čtyři barvy se střídají po stránkách. Tlumená korálová, indigo, tyrkys, zlatá.", "#D98B6A #7C78C8 #5BA8A0 #D9B86C"),
    ("zeme", "", "Čtyři barvy. Mech, okr, cihla, jezerní modrá.", "#8FBF6F #E9A23B #C9583E #4F86C6"),
    ("pastel", "", "Čtyři barvy. Tlumená růžová, aqua, levandulová, světlá citronová.", "#D9A3B8 #9CCBCF #A3B4D9 #DCD9A8"),
    ("sever", "", "Čtyři barvy. Severská modrá, šalvěj, pšeničná, tlumená červená.", "#88C0D0 #A3BE8C #EBCB8B #BF616A"),
])

section("Proměnné")
para("Do kteréhokoli textu můžete napsat proměnnou ve složených závorkách. Web ji při sestavení nahradí.")
table(["proměnná", "anglicky", "co se dosadí", "příklad"], [
    ("{rok}", "{year}", "Aktuální rok. Mění se sám, i když web dlouho nepublikujete.", "drobné: © {rok} {web}"),
    ("{web}", "{site}", "Název webu z listu site (site_name).", "text: Vítejte na {web}."),
    ("(c)", "", "Napíše se jako ©.", "drobné: (c) {rok} {web}"),
], code_first=True)

section("Obrázky a média")
for line in [
    "Obrázky nahrajte do složky media, která je ve složce klienta vedle této tabulky. Podsložky jsou v pořádku.",
    "V tabulce pak stačí název souboru: hero.jpg. Když je stejný název ve dvou podsložkách, napište cestu: galerie/hero.jpg.",
    "Web obrázky sám zmenší a převede do rychlých formátů. Nahrávejte je v plné kvalitě, nic neupravujte.",
    "Používejte JPG, PNG nebo WebP. Fotky z iPhonu (HEIC) převeďte, nebo v iPhonu nastavte Fotoaparát → Formáty → Nejkompatibilnější.",
    "Videa dávejte na YouTube nebo Vimeo a do stránky vložte odkaz.",
]:
    put(r, ["•", line])
    g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    g.row_dimensions[r].height = 28
    r += 1

section("Publikování")
for line in [
    "Až budete s úpravami hotovi, klikněte v horním menu na Web → Publikovat. Za minutu až dvě je web online.",
    "Výsledek se objeví v listu stav: kdy se publikovalo, jestli vše proběhlo a případná upozornění (třeba chybějící obrázek).",
    "Dokud nekliknete na Publikovat, návštěvníci vidí předchozí verzi. Upravovat tedy můžete v klidu.",
]:
    put(r, ["•", line])
    g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    g.row_dimensions[r].height = 28
    r += 1

section("Více jazyků")
for line in LANG_LINES:
    put(r, ["•", line])
    g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    g.row_dimensions[r].height = 28
    r += 1

section("Pravidla")
for line in [
    "id: malá písmena, bez mezer a diakritiky, slova spojená pomlčkou (t-intro, img-hero).",
    "order: počítejte po desítkách (10, 20, 30). Když později něco vložíte mezi, dáte 15 a nic se nepřečísluje.",
    "visible: no skryje stránku nebo položku, aniž byste ji mazali.",
    "Klíčové slovo s dvojtečkou vždy začíná nový kus. Nepište proto „poznámka:“ uprostřed běžné věty.",
    "Názvy sloupců (tmavé záhlaví) neměňte, web je podle nich čte.",
    "Pod patičkou každé stránky je drobný odkaz platform by filipmracek.com. Ten je součástí platformy, v tabulce se nenastavuje.",
]:
    put(r, ["•", line])
    g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    g.row_dimensions[r].height = 28
    r += 1

section("Barvy v tabulce")
put(r, ["žluté buňky", "", "Obsah stránky a texty. Tady skládáte."]); g.cell(row=r, column=1).fill = CONTENT_FILL; r += 1
put(r, ["bílé buňky", "", "Nastavení a údaje."]); r += 1
put(r, ["šedé buňky", "", "Poznámky pro vás. Na webu se neobjeví."]); g.cell(row=r, column=1).fill = NOTE_FILL; r += 1
put(r, ["roletky", "", "Sloupce jako visible, kind a focus berou jen nabízené hodnoty."]); r += 1

g.sheet_view.showGridLines = False
import sys
wb.save(sys.argv[1] if len(sys.argv) > 1 else "templates/klient-sablona.xlsx")
print("ok")
