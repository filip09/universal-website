"""Builds the control sheet template (řídicí tabulka) as .xlsx."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1C1D2E")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
AUTO_FILL = PatternFill("solid", fgColor="EEF6F3")
AUTO_FONT = Font(name=FONT, size=10, color="2F6F5E")
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="7A7890")
LINE = Border(bottom=Side(style="thin", color="E2E0EA"))
TOP = Alignment(vertical="top", wrap_text=True)

wb = Workbook()
g = wb.active
g.title = "Návod"
ws = wb.create_sheet("klienti", 0)

cols = [("client", 14), ("name", 22), ("folder", 46), ("domain", 24), ("active", 8),
        ("status", 26), ("last_publish", 20), ("url", 34), ("notes", 34)]
for c, (h, w) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    ws.column_dimensions[cell.column_letter].width = w
ws.row_dimensions[1].height = 24
example = ["ukazka", "Ukázkový web", "https://drive.google.com/drive/folders/VLOZTE-ODKAZ-NA-SLOZKU", "", "no", "", "", "", "Ukázkový řádek. Smažte ho nebo přepište."]
for c, v in enumerate(example, 1):
    cell = ws.cell(row=2, column=c, value=v)
    cell.font, cell.alignment, cell.border = (BOLD if c == 1 else BODY), TOP, LINE
for r in range(2, 300):
    for c in (6, 7, 8):
        cell = ws.cell(row=r, column=c)
        cell.fill = AUTO_FILL
        cell.font = AUTO_FONT
dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("E2:E300")
ws.freeze_panes = "B2"
ws.sheet_properties.tabColor = "F0A33C"

for col, w in zip("ABCD", (22, 18, 44, 56)):
    g.column_dimensions[col].width = w
g.sheet_properties.tabColor = "1C1D2E"
r = 1
g.cell(row=r, column=1, value="Řídicí tabulka · všechny weby na jednom místě").font = Font(name=FONT, size=18, bold=True)
r += 2


def section(title):
    global r
    g.cell(row=r, column=1, value=title).font = Font(name=FONT, size=13, bold=True)
    r += 1


def table(head, rows):
    global r
    for c, v in enumerate(head, 1):
        cell = g.cell(row=r, column=c, value=v)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    r += 1
    for row in rows:
        for c, v in enumerate(row, 1):
            cell = g.cell(row=r, column=c, value=v)
            cell.font, cell.alignment, cell.border = BODY, TOP, LINE
        r += 1
    r += 1


def bullets(lines):
    global r
    for line in lines:
        g.cell(row=r, column=1, value="•").font = BODY
        cell = g.cell(row=r, column=2, value=line)
        cell.font, cell.alignment = BODY, TOP
        g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        g.row_dimensions[r].height = 28
        r += 1
    r += 1


section("Sloupce v listu klienti")
table(["sloupec", "kdo vyplňuje", "co tam patří", "příklad"], [
    ("client", "vy", "Krátké id webu. Malá písmena, bez mezer a diakritiky. Neměňte ho, podle něj se jmenuje projekt na Cloudflare.", "studio-x"),
    ("name", "vy", "Název pro vás.", "Studio X"),
    ("folder", "vy", "Odkaz na složku klienta na Disku. Ve složce je tabulka webu a složka media.", "https://drive.google.com/drive/folders/…"),
    ("domain", "vy", "Doména webu, pokud ji klient má. Bez https://.", "www.studiox.cz"),
    ("active", "vy", "no vypne publikování tohoto webu. Web zůstane online v poslední verzi.", "yes"),
    ("status", "web sám", "Stav posledního publikování: online, upozornění, nebo chyba a proč.", "online"),
    ("last_publish", "web sám", "Kdy se naposledy publikovalo.", "22. 10. 2026 23:41"),
    ("url", "web sám", "Adresa webu. Dokud doména není připojená, je to adresa …pages.dev.", "https://uw-studio-x.pages.dev"),
    ("notes", "vy", "Poznámky pro vás. Web je nečte.", ""),
])

section("Menu Weby (nahoře v liště)")
table(["položka", "", "co udělá", ""], [
    ("Založit nového klienta…", "", "Zeptá se na id a název, založí složku klienta, zkopíruje do ní šablonu tabulky, vytvoří složku media, nasdílí vše robotovi a přidá řádek sem.", ""),
    ("Publikovat označený řádek", "", "Publikuje web klienta, na jehož řádku stojíte.", ""),
    ("Publikovat všechny weby", "", "Publikuje všechny weby s active = yes. Hodí se po vylepšení platformy.", ""),
    ("Připravit označený řádek", "", "Pro klienta, kterého jste založil ručně: nasdílí složku robotovi a vloží do jeho tabulky odkaz Publikovat.", ""),
    ("Nastavení…", "", "Jednorázově: token GitHubu, repozitář, e-mail robota, šablona, nadřazená složka, adresa webové aplikace.", ""),
])

section("Barvy")
bullets([
    "Bílé buňky vyplňujete vy.",
    "Zelené sloupce (status, last_publish, url) zapisuje web sám po každém publikování. Nepřepisujte je.",
])
g.sheet_view.showGridLines = False
wb.save(sys.argv[1] if len(sys.argv) > 1 else "templates/ridici-tabulka.xlsx")
print("ok")
