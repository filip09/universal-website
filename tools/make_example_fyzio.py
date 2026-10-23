"""Example client: a fictional physiotherapist. Fills the client template
(templates/klient-sablona.xlsx) with real content and saves it as a new sheet,
keeping the guide, formatting and dropdowns of the template."""
import sys
from copy import copy
from openpyxl import load_workbook

SRC = "templates/klient-sablona.xlsx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "examples/fyzioterapie/ladne-web.xlsx"

# ---------------------------------------------------------------- pages
PAGES = [
    ("uvod", 10, "Úvod", "Tělo, které vás zase poslechne", 18, "yes",
     "obrázek: hero | full\n"
     "úvod: Jsem Adéla, fyzioterapeutka z pražských Vinohrad. Pomáhám lidem zbavit se bolesti zad, krku a kloubů. Bez spěchu, s vysvětlením a s cviky, které doma opravdu zvládnete.\n"
     "tlačítko: Rezervovat termín -> /rezervace/\n"
     "text: t-teze | large\n"
     "seznam: cisla | cards\n"
     "nadpis: S čím za mnou lidé chodí\n"
     "text: t-uvod-pomoc\n"
     "tlačítko: Všechno, s čím pomáhám -> /sluzby/\n"
     "citace: Nechci, abyste ke mně chodili roky. Chci, abyste mě brzy nepotřebovali.\n"
     "autor: Adéla Horáková\n"
     "nadpis: Dvě nejčastější otázky\n"
     "seznam: faq-uvod | faq\n"
     "text: Další odpovědi najdete v [otázkách](/otazky/).",
     "Fyzioterapie na Vinohradech. Bolest zad, krku a kloubů, návrat po úrazu, sportovci i těhotné. Hodina jen pro vás.",
     "hero", "Úvodní stránka. Fotka přes celou šířku, hlavní sdělení a cesta k rezervaci."),
    ("sluzby", 20, "S čím pomáhám", "S čím vám pomůžu", 168, "yes",
     "úvod: Pracuji s lidmi od 15 do 85 let. Ať jde o bolest, návrat po zranění, nebo jen chcete, aby vás tělo přestalo omezovat.\n"
     "seznam: sluzby | cards\n"
     "poznámka: S čerstvým úrazem, horečkou nebo náhlou ztrátou citlivosti jděte nejdřív k lékaři. Pak se ráda postarám o návrat do pohybu.\n"
     "tlačítko: Vybrat termín -> /rezervace/",
     "Bolest zad, krku a ramen, rehabilitace po úrazu a operaci, sportovci, těhotenství a po porodu, prevence pro kancelář.",
     "sluzba-zada", "Šest karet s fotkou, obsah je v listu items (skupina sluzby)."),
    ("prubeh", 30, "Průběh", "Jak to u mě probíhá", 205, "yes",
     "úvod: Žádný běžící pás a střídání lůžek. Hodina patří jen vám.\n"
     "seznam: prubeh | steps\n"
     "obrázek: prvni-navsteva\n"
     "nadpis: Ceník\n"
     "seznam: cenik | cards\n"
     "drobné: Platba kartou nebo hotově na místě. Zrušení zdarma nejpozději 24 hodin předem.\n"
     "tlačítko: Rezervovat první návštěvu -> /rezervace/",
     "Co čekat od první návštěvy, jak terapie probíhá a kolik stojí.",
     "prvni-navsteva", "Kroky se číslují samy podle sloupce order v items."),
    ("otazky", 40, "Otázky", "Na co se lidé ptají", 42, "yes",
     "úvod: Odpovědi na to, co mi klienti píšou nejčastěji. Když tu svou otázku nenajdete, napište mi.\n"
     "seznam: faq | faq\n"
     "tlačítko: Zeptat se Adély -> /rezervace/",
     "Potřebuji doporučení? Bude to bolet? Hradí to pojišťovna? Odpovědi na nejčastější otázky.",
     "", "Rozbalovací otázky: styl faq. Otázka je title, odpověď text v items."),
    ("o-mne", 50, "O mně", "Ahoj, jsem Adéla", 330, "yes",
     "obrázek: portret\n"
     "úvod: Fyzioterapii se věnuji dvanáct let. Začínala jsem v nemocnici na rehabilitaci po operacích, dnes mám vlastní malé studio na Vinohradech.\n"
     "text: t-o-mne\n"
     "nadpis: Vzdělání a kurzy\n"
     "seznam: vzdelani | bullets\n"
     "nadpis: Studio\n"
     "text: Malá světlá místnost ve druhém patře, kde nikdo nečeká na chodbě. Čaj, ticho a hodina jen pro vás.\n"
     "galerie: studio | grid",
     "Mgr. Adéla Horáková, fyzioterapeutka s dvanáctiletou praxí. Kdo jsem, jak pracuji a jak vypadá studio.",
     "portret", ""),
    ("rezervace", 60, "Rezervace", "Pojďme se domluvit", 265, "yes",
     "úvod: Vyberte si termín, nebo mi napište, s čím přicházíte. Odpovídám do druhého dne.\n"
     "vložit: rezervacni-formular\n"
     "nadpis: Kontakt\n"
     "text: t-kontakt\n"
     "poznámka: Ordinační hodiny: pondělí až čtvrtek 8–19, pátek 8–13. Studio je ve druhém patře bez výtahu. Když vám schody dělají potíže, dejte vědět a domluvíme se jinak.",
     "Rezervace termínu a kontakt na fyzioterapeutku Adélu Horákovou, Praha 2 – Vinohrady.",
     "", "Formulář je zatím zástupný (list embeds). Kontakty jsou smyšlené, před spuštěním je přepište."),
]

# ---------------------------------------------------------------- texts
TEXTS = [
    ("t-teze", "text: Bolest není vaše povaha. Je to informace, a s informací se dá pracovat.", "Velké sdělení na úvodu (styl large)."),
    ("t-uvod-pomoc",
     "text: Bolavá záda od sezení. Krk, který nedovolí otočit hlavu. Koleno po lyžích. Rameno, se kterým se špatně spí.\n\n"
     "Většina potíží má společný kořen: tělo si zvyklo na pohyb, který mu neslouží. Najdeme ho a naučíme ho jinak.", ""),
    ("t-o-mne",
     "text: Na fyzioterapii mě baví, že každé tělo je jiný příběh. Někdo potřebuje uvolnit, jiný posílit a většina lidí hlavně pochopit, co se v těle děje. Proto hodně vysvětluji a kreslím.\n\n"
     "Ve volném čase lezu, běhám po Stromovce a učím se hrát na violoncello. To poslední jde nejpomaleji.", ""),
    ("t-kontakt",
     "text: **Telefon:** +420 600 000 000\n**E-mail:** ahoj@fyzio-ladne.example\n**Kde:** Studio Ladně, Praha 2 – Vinohrady\n\n"
     "Tramvají na zastávku Jiřího z Poděbrad, odtud 3 minuty pěšky.",
     "Kontakty jsou smyšlené (ukázkový web). Jeden řádek = Ctrl+Enter, prázdný řádek = nový odstavec."),
]

# ---------------------------------------------------------------- items
def group(name, rows, step=10):
    return [(name, (i + 1) * step, *r) for i, r in enumerate(rows)]

FAQ = [
    ("Potřebuji doporučení od lékaře?", "Ne, můžete přijít rovnou. Pokud máte zprávy z vyšetření, rentgen nebo magnetickou rezonanci, vezměte je s sebou. Ušetří nám to čas."),
    ("Bude to bolet?", "Terapie může být místy nepříjemná, ale nikdy nepracuji přes bolest, kterou nechcete. Celou dobu se ptám a tempo určujete vy."),
    ("Co si mám vzít s sebou?", "Pohodlné oblečení, ve kterém se dá hýbat, třeba legíny nebo kraťasy a tričko. Ručník, podložku i vodu mám ve studiu."),
    ("Kolik návštěv budu potřebovat?", "Většině lidí se výrazně uleví po 3 až 5 setkáních. U dlouhodobých potíží nebo po operaci jich bývá víc. Odhad vám řeknu po prvním vyšetření."),
    ("Jak rychle se mi uleví?", "Často už po první hodině, protože uvolníme to nejvíc přetížené. Aby úleva vydržela, je potřeba změnit pohybové návyky, a to trvá pár týdnů."),
    ("Hradí to zdravotní pojišťovna?", "Ne. Nejsem smluvní fyzioterapeutka, a proto mám na každého celou hodinu. Péči ale často proplatí zaměstnanecké benefity. Doklad vám ráda vystavím."),
    ("Můžu přijít v těhotenství?", "Ano, ve kterémkoli trimestru, pokud vám to nezakázal lékař. Pracuji v poloze, která je pro vás pohodlná, a vynechávám techniky, které se v těhotenství nehodí."),
    ("Pracujete i s dětmi?", "S dospívajícími od 15 let ano, třeba se skoliózou nebo sportovním přetížením. Na první návštěvu přicházejí s rodičem."),
    ("Co když musím termín zrušit?", "Napište nebo zavolejte nejpozději 24 hodin předem a nic neplatíte. Při pozdějším zrušení účtuji polovinu ceny, protože termín už nestihnu nabídnout jinému."),
    ("Dostanu cviky domů?", "Vždycky. Tři cviky, které spolu natrénujeme na místě, a k nim krátké video do mobilu, abyste si nemuseli nic pamatovat."),
]

ITEMS = []
ITEMS += group("cisla", [
    ("12 let praxe", "S klienty od 15 do 85 let, od kancelářských zad po návrat na maraton.", "", "", "yes"),
    ("60 minut jen pro vás", "Žádné střídání lůžek ani čekání na chodbě.", "", "", "yes"),
    ("3 cviky domů", "Ne dvacet, tři. Ty, které opravdu budete dělat.", "", "", "yes"),
])
ITEMS += group("sluzby", [
    ("Bolest zad a beder", "Sedavá práce, zvedání dětí, dlouhé řízení. Uvolníme, co je přetížené, a posílíme, co zahálí.", "sluzba-zada", "", "yes"),
    ("Krk, ramena a hlava", "Ztuhlá šíje, bolest hlavy z napětí, rameno, které nejde zvednout nad hlavu.", "sluzba-krk", "", "yes"),
    ("Po úrazu a operaci", "Kotník, koleno, kyčel, rameno. Bezpečný návrat k pohybu krok za krokem, ve spolupráci s vaším lékařem.", "sluzba-uraz", "", "yes"),
    ("Sportovci", "Běžci, lezci, tenisté. Řešíme přetížení dřív, než z něj bude zranění.", "sluzba-sport", "", "yes"),
    ("Těhotenství a po porodu", "Bolest zad v těhotenství, pánevní dno, rozestup břišních svalů. Jemně a vaším tempem.", "sluzba-tehotenstvi", "", "yes"),
    ("Kancelářské tělo", "Prevence pro ty, kdo sedí osm hodin denně. Úprava pracovního místa a cviky, které se vejdou mezi dva e-maily.", "sluzba-kancelar", "", "yes"),
])
ITEMS += group("prubeh", [
    ("Povídání", "Zeptám se na vaše potíže, práci, sport i spánek. Bolest má většinou příběh a ten nás dovede k příčině.", "", "", "yes"),
    ("Vyšetření", "Podívám se, jak stojíte, chodíte a dýcháte. Otestuji pohyblivost a sílu a hned vysvětlím, co vidím.", "", "", "yes"),
    ("Terapie", "Uvolnění měkkých tkání, mobilizace kloubů, práce s dechem. Podle toho, co vaše tělo zrovna potřebuje.", "", "", "yes"),
    ("Cviky domů", "Odnesete si tři cviky a k nim krátké video do mobilu.", "", "", "yes"),
    ("Kontrola", "Příště zjistíme, co zabralo, a posuneme se dál. Většině lidí stačí 3 až 5 setkání.", "", "", "yes"),
])
ITEMS += group("cenik", [
    ("Vstupní vyšetření a terapie · 60 min", "**1 600 Kč.** Vyšetření, první ošetření a plán na další týdny.", "", "", "yes"),
    ("Navazující terapie · 50 min", "**1 300 Kč.**", "", "", "yes"),
    ("Balíček 5 terapií", "**6 000 Kč.** Ušetříte 500 Kč, platí půl roku.", "", "", "yes"),
    ("Online konzultace · 30 min", "**700 Kč.** Kontrola cviků nebo pracovního místa přes video.", "", "", "yes"),
])
ITEMS += group("faq", [(q, a, "", "", "yes") for q, a in FAQ])
ITEMS += group("faq-uvod", [(q, a, "", "", "yes") for q, a in FAQ if q.startswith(("Bude to", "Potřebuji"))])
ITEMS += group("vzdelani", [
    ("Magisterské studium fyzioterapie.", "Praha.", "", "", "yes"),
    ("Dynamická neuromuskulární stabilizace (DNS).", "Kurzy A až C.", "", "", "yes"),
    ("Fyzioterapie pánevního dna.", "Péče v těhotenství a po porodu.", "", "", "yes"),
    ("Sportovní fyzioterapie a prevence zranění.", "", "", "", "yes"),
    ("Kinesiotaping.", "", "", "", "yes"),
])
ITEMS += group("studio", [("", "", f"studio-0{n}", "", "yes") for n in range(1, 7)])

# ---------------------------------------------------------------- media
# id, kind, source, alt, caption, credit, focus, notes (= what the photo should show)
L32 = "Na šířku, aspoň 2400 px. "
CARD = "Na šířku, aspoň 1600 px. Na webu se ořízne na poměr 16:9, hlavní motiv dejte doprostřed."
TILE = "Na výšku, aspoň 1200 px. Na webu se ořízne na poměr 4:5."
MEDIA = [
    ("hero", "image", "hero.jpg", "Adéla uvolňuje klientce ramena ve světlém studiu", "", "", "center",
     L32 + "Hlavní fotka přes celou šířku úvodu, slouží i jako náhled při sdílení odkazu. Adéla při práci s klientkou: ruce na rameni nebo lopatce, obě uvolněné, klidně úsměv. Světlé studio, denní světlo. Důležité dejte doprostřed, na telefonu se okraje ořežou."),
    ("portret", "image", "portret.jpg", "Adéla Horáková, fyzioterapeutka", "", "", "top",
     "Na výšku 4:5, aspoň 1600 px na šířku. Portrét do pasu na stránku O mně. Přirozený úsměv, pohled do objektivu, tričko nebo civilní oblečení, za ní rozostřené studio."),
    ("ikona", "image", "ikona.png", "Ladně", "", "", "center",
     "Čtverec 512 × 512 px, PNG. Logo nebo monogram do záložky prohlížeče. Jednoduchý tvar, který je čitelný i v 32 px."),
    ("prvni-navsteva", "image", "prvni-navsteva.jpg", "Vyšetření držení těla při první návštěvě", "První návštěva začíná rozhovorem a vyšetřením.", "", "center",
     "Na šířku 3:2, aspoň 1600 px. Vyšetření: Adéla pozoruje, jak klient stojí, nebo testuje pohyblivost kyčle či ramene. Stránka Průběh."),
    ("sluzba-zada", "image", "sluzby/zada.jpg", "Ošetření bolavých beder", "", "", "center",
     CARD + " Klient leží na břiše, ruce terapeutky na dolní části zad."),
    ("sluzba-krk", "image", "sluzby/krk.jpg", "Uvolnění šíje a ramen", "", "", "center",
     CARD + " Uvolnění krku: klient leží na zádech nebo sedí, ruce terapeutky na šíji."),
    ("sluzba-uraz", "image", "sluzby/uraz.jpg", "Cvičení s kolenem po zranění", "", "", "center",
     CARD + " Rehabilitace: práce s kolenem nebo kotníkem, případně cvik s odporovou gumou."),
    ("sluzba-sport", "image", "sluzby/sport.jpg", "Sportovec při stabilizačním cvičení", "", "", "center",
     CARD + " Sportovec při cviku: dřep na jedné noze, balanční podložka nebo protažení po běhu."),
    ("sluzba-tehotenstvi", "image", "sluzby/tehotenstvi.jpg", "Těhotná žena cvičí na velkém míči", "", "", "center",
     CARD + " Těhotná žena při jemném cvičení na velkém míči nebo podložce, klidná atmosféra."),
    ("sluzba-kancelar", "image", "sluzby/kancelar.jpg", "Protažení u pracovního stolu", "", "", "center",
     CARD + " Člověk u stolu s notebookem dělá jednoduchý protahovací cvik, nebo Adéla upravuje výšku monitoru."),
    ("studio-01", "image", "studio/studio-01.jpg", "Pohled do studia s terapeutickým lehátkem", "", "", "center", TILE + " Celkový pohled do studia s lehátkem."),
    ("studio-02", "image", "studio/studio-02.jpg", "Ruce terapeutky při ošetření", "", "", "center", TILE + " Detail rukou při terapii."),
    ("studio-03", "image", "studio/studio-03.jpg", "Cvičební pomůcky na polici", "", "", "center", TILE + " Pomůcky: míče, gumy, válce, podložky."),
    ("studio-04", "image", "studio/studio-04.jpg", "Okno studia s rostlinami", "", "", "center", TILE + " Okno, denní světlo, rostliny."),
    ("studio-05", "image", "studio/studio-05.jpg", "Adéla ukazuje klientovi cvik", "", "", "center", TILE + " Adéla předvádí cvik klientovi."),
    ("studio-06", "image", "studio/studio-06.jpg", "Čaj a klidné místo na odpočinek", "", "", "center", TILE + " Detail atmosféry: čaj, recepce nebo cedulka na dveřích studia."),
]

# ---------------------------------------------------------------- embeds
FIELD = "padding:12px 14px;border-radius:10px;border:1px solid rgba(255,255,255,.14);background:rgba(255,255,255,.04);color:var(--muted);font-size:.92em"
FORM = f"""<div style="border:1px dashed hsla(var(--h),70%,75%,.45);border-radius:16px;padding:22px;display:grid;gap:12px">
<div style="font-family:var(--f-mono);font-size:11px;letter-spacing:.08em;text-transform:uppercase;color:hsl(var(--h),70%,78%)">Online rezervace · připravujeme</div>
<div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:10px">
<div style="{FIELD}">Jméno a příjmení</div><div style="{FIELD}">Telefon nebo e-mail</div>
<div style="{FIELD}">Preferovaný den a čas</div><div style="{FIELD}">Typ návštěvy</div>
</div>
<div style="{FIELD};min-height:84px">S čím přicházíte?</div>
<div style="display:flex;flex-wrap:wrap;align-items:center;gap:12px 18px"><span style="padding:11px 22px;border-radius:999px;background:hsla(var(--h),70%,78%,.35);color:#fff;font-weight:600">Odeslat</span><span style="color:var(--muted);font-size:.9em">Do té doby volejte nebo pište, kontakt je níže.</span></div>
</div>"""
EMBEDS = [("rezervacni-formular", FORM, "", "Zástupný blok. Až bude rezervační systém, nahradíme ho skutečným formulářem.")]

# ---------------------------------------------------------------- settings
SITE = {
    "site_name": "Ladně · fyzioterapie",
    "language": "cs",
    "domain": "",
    "home_page": "uvod",
    "description": "Fyzioterapie Adély Horákové na Vinohradech. Bolest zad, krku a kloubů, návrat po úrazu, sportovci i těhotné.",
    "share_image": "hero",
    "favicon": "ikona",
    "footer": "drobné: © {rok} {web} · Mgr. Adéla Horáková · Praha 2 – Vinohrady",
}
STYLES = {"font_display": "Fraunces", "font_body": "Figtree", "glass_opacity": 0.62, "text_size": 17}

# ---------------------------------------------------------------- write
wb = load_workbook(SRC)


def fill(name, rows):
    ws = wb[name]
    width = ws.max_column
    styles = [copy(ws.cell(row=2, column=c)._style) for c in range(1, width + 1)]
    for r in range(2, max(ws.max_row, len(rows) + 1) + 1):
        for c in range(1, width + 1):
            ws.cell(row=r, column=c).value = None
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v if v != "" else None)
            cell._style = copy(styles[c - 1])


def settings(name, values):
    ws = wb[name]
    found = set()
    for r in range(2, ws.max_row + 1):
        key = str(ws.cell(row=r, column=1).value or "").strip()
        if key in values:
            ws.cell(row=r, column=2).value = values[key] if values[key] != "" else None
            found.add(key)
    missing = set(values) - found
    assert not missing, f"{name}: missing settings {missing}"


fill("pages", PAGES)
fill("texts", TEXTS)
fill("items", ITEMS)
fill("media", MEDIA)
fill("embeds", EMBEDS)
settings("site", SITE)
settings("styles", STYLES)
wb["Návod"].cell(row=1, column=1).value = "Ladně · fyzioterapie · návod k tabulce"
wb.active = wb.sheetnames.index("pages")
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == "pages"
wb.save(OUT)
print("ok", OUT, len(PAGES), "pages,", len(ITEMS), "items,", len(MEDIA), "images")
